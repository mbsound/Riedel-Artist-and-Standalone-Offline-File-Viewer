#!/usr/bin/env python3
"""
bol_faithful.py — 1:1 dead accurate parser of Riedel Bolero .bol save files.
Direct, faithful port of libRadon.so's (v3.4.1) own CombinedNetConfig::unpackFromSaved routines.

ALL 9 SECTIONS FULLY VALIDATED & EXACT-CONSUMPTION GATED ACROSS ALL 8 SAMPLE FILES:
 1. network        (v10, radon::NetSettings::unpackDataFromSaved)
 2. partylines     (v2,  radon::Partylines::unpackDataFromSaved)
 3. profiles       (v17, radon::Profiles::unpackDataFromSaved)
 4. beltpacks      (v19, radon::RegisteredBPs::unpackDataFromSaved)
 5. antennas       (v1,  radon::priorityNodes::PriorityNodes::unpackDataFromSaved)
 6. audio_filters  (v1,  radon::AudioFilters::unpackDataFromSaved)
 7. audio_devices  (v6,  radon::IODeviceConfig::unpackDataFromSaved)
 8. audio_channels (v6,  radon::AudioChannels::unpackDataFromSaved)
 9. gpio           (v60000, radon::TriggerConfig::unpackDataFromSaved)

Correctness gate: cursor == section end (zero residual bytes) for every single section on every file.
Enums and layouts are decompiled directly from libRadon.so (ARM ELF).
"""
import struct, zlib, sys, glob

# ── enums (firmware-confirmed) ──────────────────────────────────────────────
KEY_FUNCTIONS = {0:"None",1:"Talk",2:"Talk+AlwaysListen",3:"Talk&Listen",4:"Listen",
    5:"Monitor",6:"MonitorSelect",7:"NotificationBeep",8:"NotificationBeepSelect",9:"Reply",
    10:"MenuShortcut",11:"Toggle",12:"MonitorTrigger",13:"SetTrigger",14:"VolumeIncrease",
    15:"VolumeDecrease",16:"MuteMic",17:"Control",18:"ControlWithDestination"}
KEY_MODES = {0:"Momentary",1:"Latching",2:"Auto",3:"On only",4:"Off only"}
KEY_PRIORITIES = {0:"Standard",1:"High",2:"Low"}
# web GUI CONFIG enums (3.4.1 app-*.js)
KEY_GROUPS = {0:"Off",1:"1",2:"2",3:"3",4:"4",5:"5"}
MUTED_KEY_PRESS_ACTIONS = {0:"Keep Mute State",1:"Unmute",2:"Momentary Unmute"}
DOUBLE_PRESS_MODES = {0:"Off",1:"Monitor",2:"Forced Listen"}
MIC_TYPES = {0:"Auto",1:"Dynamic Detect",2:"Electret Detect",3:"Dynamic",4:"Electret",5:"RSM detect",6:"RSM"}
DISPLAY_MODES = {0:"Standard",1:"Alternative",2:"Standard Flip",3:"Alternative Flip"}
LANGUAGES = {0:"English",1:"Deutsch",2:"Chinese"}
BRIGHTNESS_MODES = {0:"Off",1:"Low",2:"Medium",3:"High",4:"Custom"}
BLUETOOTH_MODES = {0:"Local",1:"Public"}
BLUETOOTH_STATES = {0:"Off",1:"Connect to Headset",2:"Connect to Mobile/PC",4:"Connect to SmartWatch"}
BP_PRIORITIES = {0:"Not Allowed",1:"Low",2:"Medium",3:"High"}     # BP_DEFAULT_PRIORITY_TYPES
AUDIO_PORT_TYPE = {0:"None",1:"BP(P2P)",2:"PL",3:"AudioChannelBiDir",4:"KeyParameter",
    5:"AudioChannelInput",6:"AudioChannelOutput",7:"TriggerInput",8:"TriggerOutput",9:"TriggerVirtual"}
AUDIO_CHANNEL_TYPE = {0:"Unknown",1:"LegacyNSA",2:"PQ_Analog",3:"PQ_Partyline",9:"UniversalAutomatic",
    10:"UniversalAnalog",11:"UniversalDigitalLeft",12:"UniversalDigitalRight",13:"UniversalDigitalBoth",
    14:"USBAudioLeft",15:"USBAudioRight",16:"USBAudioBoth",17:"Headphone",18:"TwoWireFirst",19:"TwoWireSecond"}
IO_DEVICE_TYPE = {0:"Unknown",1:"NSA-002A",2:"PunQtum",3:"NSA-003A",4:"NSA-004A",5:"NSA-005A",
    6:"NSA-006A",7:"NSA-007A",8:"NSA-008A",9:"NSA-009A",10:"NSA-010C"}
TRIGGER_MODES = {0:"normal",1:"forceOn",2:"forceOff"}


class InputStream:
    def __init__(self, data, start=0, end=None):
        self.d = data; self.p = start; self.end = len(data) if end is None else end
    def _need(self, n):
        if self.p + n > self.end:
            raise EOFError(f"read past section end: p={self.p} need={n} end={self.end}")
    def u8(self):  self._need(1); v=self.d[self.p]; self.p+=1; return v
    def i8(self):  v=self.u8(); return v-256 if v>=128 else v
    def boolean(self): return self.u8()!=0
    def u16(self): self._need(2); v=struct.unpack_from(">H",self.d,self.p)[0]; self.p+=2; return v
    def i16(self): v=self.u16(); return v-0x10000 if v>=0x8000 else v
    def u32(self): self._need(4); v=struct.unpack_from(">I",self.d,self.p)[0]; self.p+=4; return v
    def read(self, n): self._need(n); v=self.d[self.p:self.p+n]; self.p+=n; return v
    def string(self):
        n=self.u8(); return self.read(n).decode("utf-8","replace")

# ── leaf unpackers (mirror radon::X::unpack) ────────────────────────────────
def up_audio_port_id(s):
    t=s.u8()
    return {"type":t,"type_name":AUDIO_PORT_TYPE.get(t,f"?{t}"),
            "id": s.u16() if t!=0 else None}

def up_function_specific_parameter(s):
    return {"a":s.i8(),"b":s.u32()}

def up_bpkey(s):
    """radon::BPKey::unpack — fixed 6 bytes + AudioPortId + FSP(if function==7)."""
    fn=s.u8(); mode=s.u8(); prio=s.u8(); destflags=s.u8(); muted=s.u8(); dpm=s.u8()
    target=up_audio_port_id(s)
    fsp=up_function_specific_parameter(s) if fn==7 else None
    return {"function":fn,"function_name":KEY_FUNCTIONS.get(fn,f"?{fn}"),
            "mode":mode,"mode_name":KEY_MODES.get(mode,f"?{mode}"),
            "priority":prio,"priority_name":KEY_PRIORITIES.get(prio,f"?{prio}"),
            # JSON keys: priority, keyGroup, mutedKeyPressAction, doublePressMode (serializeBPConfig)
            "key_group":destflags,"key_group_name":KEY_GROUPS.get(destflags, destflags),
            "muted_key_press_action":muted,"muted_key_press_action_name":MUTED_KEY_PRESS_ACTIONS.get(muted, muted),
            "double_press_mode":dpm,"double_press_mode_name":DOUBLE_PRESS_MODES.get(dpm, dpm),
            "target":target,"fsp":fsp}

# Field names below are libRadon's own JSON keys (JsonSerializer::serialize*, strings resolved from libRadon.so
# with tools/radon_json_keys.py), matched to the member offsets each unpack() writes (tools/radon_keymap.py).

def up_single_function(s):
    """SingleFunction: function, priority, audioPortId (+ function-specific parameter for function 7)."""
    fn=s.u8(); prio=s.u8(); target=up_audio_port_id(s)
    fsp=up_function_specific_parameter(s) if fn==7 else None
    return {"function":fn,"function_name":KEY_FUNCTIONS.get(fn,f"?{fn}"),
            "priority":prio,"priority_name":KEY_PRIORITIES.get(prio,f"?{prio}"),"target":target,"fsp":fsp}

def up_trigger_function(s):
    """TriggerFunction: the trigger that fires it, then a SingleFunction."""
    return {"port":up_audio_port_id(s),"function":up_single_function(s)}

def up_external_bpkey(s):
    """ExternalBPKey (JSON externalKeys / remoteKeys): u8 +4, u16 id +8."""
    return {"a":s.u8(),"id":s.u16()}

# AudioPortsList entry usage bits: AudioPortsList::createUsage / createFunctionTypeUsage and the web GUI's
# AUDIO_PORT_USAGE {KEY 0, ALWAYS_ON 1, REPLY 2, TRIGGER 3, ON_CALL 4, ON_SILENT_CALL 5, ON_VOX 6} as bits;
# 0x80 is set for a Talk+AlwaysListen function (its always-listen half).
AUDIO_PORT_USAGE = {0x01: "Key", 0x02: "Always-On", 0x04: "Reply", 0x08: "Trigger", 0x10: "On-Talk",
                    0x20: "On-Notification/Beep", 0x40: "On-VOX", 0x80: "Always-Listen"}

def up_audio_ports_list(s):
    """AudioPortsList entries (JSON audioPorts): id, volume, mute, usage bits, then one count per usage bit."""
    n=s.u8(); entries=[]
    for _ in range(n):
        pid=up_audio_port_id(s); vol=s.i8(); mute=s.boolean(); usage=s.u8()
        counts=[s.u8() for bit in range(8) if usage & (1<<bit)]
        entries.append({"port":pid,"volume":vol,"mute":mute,"usage":usage,
                        "usage_names":[n for b, n in AUDIO_PORT_USAGE.items() if usage & b],
                        "usage_counts":counts})
    return entries

# BPVolumeData: 17 bytes, BPVolumeData::get(i) order as used by JsonSerializer::serializeBPConfig
BP_VOLUME_NAMES = ["headset_volume", "speaker_volume", "sidetone_volume", "headset_mic_input_gain",
                   "internal_mic_input_gain", "aux_input_gain", "signalization_beep_volume",
                   "signalization_voice_volume", "mic_limiter", "bluetooth_volume", "headset_lower_limit",
                   "speaker_lower_limit", "priority_dim", "headset_limiter", "bluetooth_mic_input_gain",
                   "rsm_volume", "rsm_headset_delta"]
def up_bp_volume_data(s):
    raw = s.read(0x11)
    return {n: (v - 256 if v > 127 else v) for n, v in zip(BP_VOLUME_NAMES, raw)}   # signed dB values
# BPAudioFilters +0xd4..+0xe8
BP_FILTER_NAMES = ["mic_filter", "headphone_filter", "internal_mic_filter", "internal_speaker_filter",
                   "rsm_mic_filter", "rsm_speaker_filter"]
def up_bp_audio_filters(s): return {n: s.u8() for n in BP_FILTER_NAMES}
def up_bp_brightness(s):
    """BPBrightness (JSON <prefix>Brightness, BrightnessDimmed, DimTimer, OffTimer)."""
    return {"brightness":s.u8(),"dimmed":s.u8(),"dim_timer":s.u8(),"off_timer":s.u8()}
def up_bp_bluetooth(s):
    """BPBluetooth: bluetoothState, bluetoothMode, bluetoothDimLevel."""
    st=s.u8(); md=s.u8()
    return {"state":st,"state_name":BLUETOOTH_STATES.get(st, st),"mode":md,"mode_name":BLUETOOTH_MODES.get(md, md),
            "dim_level":s.i8()}
# BPSignalization: one u32, 4 bits per event (BPSignalization::get(i) = value >> 4i & 0xf); bits are the web GUI
# BPSignalization {light 1, vibrate 2, beep 4, voice 8}
SIGNALIZATION_EVENTS = ["call", "silent_call", "low_battery", "out_of_range", "key_volume"]
SIGNALIZATION_BITS = {1: "light", 2: "vibrate", 4: "beep", 8: "voice"}
def signalization_text(v):
    return ", ".join(n for b, n in SIGNALIZATION_BITS.items() if v & b) or "none"
def up_bp_signalization(s):
    raw = s.u32()
    out = {e: signalization_text(raw >> (4 * i) & 0xf) for i, e in enumerate(SIGNALIZATION_EVENTS)}
    out["raw"] = raw
    return out
def up_bp_change_rights(s): return s.u32()
def up_voxbase(s):
    """VoxBase::unpack: u8, i8, i8, u8, u16, u16, bool = JSON vox {state, onThreshold, hysteresis, delta,
    holdTime, releaseTime, noiseGate} (JsonSerializer::serializeVox reads +0, +4, +5, +6, +8, +10, +0xc)."""
    return {"state":s.u8(),"on_threshold":s.i8(),"hysteresis":s.i8(),"delta":s.u8(),
            "hold_time":s.u16(),"release_time":s.u16(),"noise_gate":s.boolean()}

VOX_STATES = {0: "Off", 1: "Standard", 2: "Adaptive"}          # web GUI VOX_STATES

def vox_text(v):
    if not v or not v.get("state"):
        return "Off"
    return (f"{VOX_STATES.get(v['state'], v['state'])}, threshold {v['on_threshold']} dB, "
            f"hold {v['hold_time']} ms, release {v['release_time']} ms"
            + (", noise gate" if v.get("noise_gate") else ""))
def up_bp_quickmenu(s):
    """BPQuickMenu (JSON quickMenuEntries): raw entry bytes."""
    n=s.u8(); return list(s.read(n))
def up_bp_priority_config(s):
    """BPPriorityConfig: defaultPriority, then priorityExceptions {nodeId, priority}."""
    d=s.u8(); n=s.u8(); items=[{"node":s.u32(),"priority":s.u8()} for _ in range(n)]
    for it in items:
        it["priority_name"] = BP_PRIORITIES.get(it["priority"], it["priority"])
    return {"default_priority":d,"default_priority_name":BP_PRIORITIES.get(d, d),"exceptions":items}

def up_artist_bp_key(s):   # ArtistBPKey: string + 6*u8 + u16 + u8
    return {"name":s.string(),"b":[s.u8() for _ in range(6)],"u16":s.u16(),"c":s.u8()}
def up_artist_bp_config(s):
    name=s.string(); vox=up_voxbase(s)
    n=s.u16(); vol_mutes=[{"key":s.u16(),"vol":s.i8(),"mute":s.boolean()} for _ in range(n)]
    keys=[up_artist_bp_key(s) for _ in range(7)]   # ctor _M_default_append(...,7)
    return {"name":name,"vox":vox,"volume_mutes":vol_mutes,"artist_keys":keys}

# counts that are pre-sized vectors (not in the stream); from ctor
N_KEYS=7; N_EXTERNAL_KEYS=4; N_ROTARIES=2

# BPConfig flag byte +0xfc (JsonSerializer::serializeBPConfig bit tests; deserializer masks agree)
BP_FLAG_BITS = {0x01: "speaker_enable", 0x02: "silent_mode", 0x04: "echo_cancellation",
                0x08: "plug_func_activate_headset", 0x10: "plug_func_activate_speaker",
                0x20: "allow_multi_registration", 0x40: "automatic_net_change", 0x80: "show_on_reply"}

def up_bpconfig(s):
    """BPConfig::unpack (a beltpack's or a profile's configuration); names are serializeBPConfig's JSON keys."""
    c={}
    c["id"]=s.i16()                                   # +0x04 profileId
    c["name"]=s.string()                              # bpName
    c["bp_number"]=s.u16()                            # +0x20 bpNumber
    c["u8_24"]=s.u8()                                 # +0x24: not in the JSON; meaning unknown
    c["audio_ports"]=up_audio_ports_list(s)           # +0x28 audioPorts
    c["keys"]=[up_bpkey(s) for _ in range(N_KEYS)]    # +0x44 keys
    c["always_on"]=[up_single_function(s) for _ in range(s.u8())]        # +0x50 alwaysOnFunctions
    c["triggers"]=[up_trigger_function(s) for _ in range(s.u8())]       # +0x60 triggerFunctions
    c["on_talk"]=[up_single_function(s) for _ in range(s.u8())]          # +0x70 onCallFunctions
    c["on_beep"]=[up_single_function(s) for _ in range(s.u8())]          # +0x80 onSilentCallFunctions
    c["on_vox"]=[up_single_function(s) for _ in range(s.u8())]           # +0x90 onVoxFunctions
    c["rotaries"]=[s.u16() for _ in range(N_ROTARIES)]                  # rotaries
    c["external_keys"]=[up_external_bpkey(s) for _ in range(N_EXTERNAL_KEYS)]   # +0xac externalKeys
    c["volumes"]=up_bp_volume_data(s)                 # +0xb8 BPVolumeData
    c["mic_type"]=s.u8()                              # +0xd0 micType
    c["mic_type_name"]=MIC_TYPES.get(c["mic_type"], c["mic_type"])
    c["audio_filters"]=up_bp_audio_filters(s)         # +0xd4..+0xe8
    c["noise_filter"]=bool(s.u8())                    # +0xec noiseFilter
    c["display_mode"]=s.u8()                          # +0xf0 displayMode
    c["display_mode_name"]=DISPLAY_MODES.get(c["display_mode"], c["display_mode"])
    c["skinny_key_volume_override"]=s.u8()            # +0xf4 skinnyKeyVolumeOverride
    c["language"]=s.u8()                              # +0xf8 language
    c["language_name"]=LANGUAGES.get(c["language"], c["language"])
    flags=s.u8()                                      # +0xfc flag byte
    c["flags"]={n: bool(flags & b) for b, n in BP_FLAG_BITS.items()}
    c["replay_recording_time"]=s.i8()                 # +0x100 replayRecordingTime
    c["replay_store_time"]=s.u8()                     # +0x104 replayStoreTime
    c["timeout_menu"]=s.u8()                          # +0x108 timeoutMenu
    c["timeout_volume"]=s.u8()                        # +0x10c timeoutVolume
    c["timeout_off"]=s.u8()                           # +0x110 timeoutOff
    c["brightness_mode"]=s.u8()                       # +0x114 brightnessMode
    c["brightness_mode_name"]=BRIGHTNESS_MODES.get(c["brightness_mode"], c["brightness_mode"])
    c["display"]=up_bp_brightness(s)                  # +0x118 display*
    c["key_brightness"]=up_bp_brightness(s)           # +0x12c key*
    c["call_led_dim"]=s.u8(); c["status_led_dim"]=s.u8()     # +0x140, +0x144
    c["bluetooth"]=up_bp_bluetooth(s)                 # +0x148 bluetoothState/Mode/DimLevel
    c["signalization"]=up_bp_signalization(s)         # +0x158 call/silentCall/lowBatt/outOfRange/keyVolume
    c["default_signalization_pattern"]=s.u16()       # +0x160 defaultSignalizationPattern
    c["quick_menu_entries"]=up_bp_quickmenu(s)        # +0x168 quickMenuEntries
    c["priority"]=up_bp_priority_config(s)            # +0x174 defaultPriority, priorityExceptions
    c["change_rights"]=up_bp_change_rights(s)         # +0x194 changeRights
    c["vox"]=up_voxbase(s); c["vad"]=up_voxbase(s)    # +0x198 vox, +0x1a8 vad
    c["reply_partyline_mode"]=s.u8()                  # +0x1b8 partyLineReplyMode
    c["artist"]=up_artist_bp_config(s)                # +0x1bc ArtistBPConfig
    return c

def up_profile(s):
    """Profile::unpack: name, BPConfig, +0x25c appendIdToDefaultName, +0x25d updateName, +0x260 timestamp."""
    return {"name":s.string(),"config":up_bpconfig(s),
            "append_id_to_default_name":s.boolean(),"update_name":s.boolean(),"timestamp":s.u32()}

OVERRIDE_MODES = {0: "Off", 1: "Active", 2: "Active + written to beltpack config"}

# ── section 0: network (radon::NetSettings::unpackDataFromSaved / NetSettingsData::unpack) ──
def up_net_settings_section(s):
    net_name = s.string()
    system_mode = s.u8(); admin_pin = s.i16(); ota_pin = s.i16()
    audio_multicast_group = list(s.read(4)); multicast_ttl = s.u8(); time_source = s.u8()
    net_label = s.string()
    time_offset = s.u32()
    time_format = s.u8(); date_format = s.u8()
    radio_power = s.u8(); radio_power_2g4 = s.u8(); radio_retransmission_limit = s.u8()
    frequency_hopping_mode = s.u8(); radio_flags = s.u8()
    # RegistrationMode::unpack: +4 flags, +8 timeout, +0xc profileId (JSON registrationMode: registrationEnabled =
    # flags != 0, otaEnabled bit 0, nfcEnabled bit 1, chargerEnabled bit 2, timeout, profileId)
    rm_flags = s.u8(); rm_timeout = s.i8(); rm_profile = s.i16()
    rm = {"flags": rm_flags, "registration_enabled": rm_flags != 0, "ota_enabled": bool(rm_flags & 1),
          "nfc_enabled": bool(rm_flags & 2), "charger_enabled": bool(rm_flags & 4),
          "timeout": rm_timeout, "profile_id": rm_profile}
    broadcast_mode = s.u8(); enc = s.read(32)
    ptp_domain = s.u8(); debug_flags = s.u32()
    dscp = [s.u8() for _ in range(3)]
    ptp_mode = s.u8()
    # ArtistNetSettings (JSON artistNetSettings): networkId, clusterId, mcAnnounceIp, mcAnnouncePort
    ans = {"network_id": s.u32(), "cluster_id": s.u32(), "mc_announce_ip": ".".join(str(b) for b in s.read(4)),
           "mc_announce_port": s.u16()}
    bp_mon_threshold = s.i16()
    bp_mon_blocked_rssi = s.i8(); bp_mon_interfered_rssi = s.i8()
    dect_scanner_blocked_rssi = s.i8(); dect_scanner_interfered_rssi = s.i8()
    term_cnt = s.u8()
    terms = [s.u16() for _ in range(term_cnt)]
    # BPOverrideSettings: 4 polymorphic entries in fixed order. Each starts with the override mode
    # (deserializeBPOverrideMode: 0 off, 1 overrideActive, 2 overrideActive + overwriteBPconfig); field names are
    # the web GUI's BpOverrideSettings.
    # 0: Bluetooth (4 bytes)
    bt_mode = s.u8(); bt = up_bp_bluetooth(s)
    # 1: Brightness (12 bytes)
    br_mode = s.u8(); br_sub = s.u8()
    br1 = up_bp_brightness(s); br2 = up_bp_brightness(s)
    br_i1 = s.i8(); br_i2 = s.i8()
    # 2: Notification (7 bytes)
    notif_mode = s.u8(); notif_sig = s.u32()
    notif_i1 = s.i8(); notif_i2 = s.i8()
    # 3: Speaker (2 bytes)
    spk_mode = s.u8(); spk_enabled = s.boolean()
    sig = {e: signalization_text(notif_sig >> (4 * i) & 0xf) for i, e in enumerate(SIGNALIZATION_EVENTS)}

    return {
        "show_name": net_name, "label": net_label, "system_mode": system_mode,
        "admin_pin": admin_pin, "ota_pin": ota_pin,
        "audio_multicast_group": audio_multicast_group, "multicast_ttl": multicast_ttl,
        "time_source": time_source, "time_offset": time_offset,
        "time_format": time_format, "date_format": date_format,
        "radio_power": radio_power, "radio_power_2g4": radio_power_2g4,
        "radio_retransmission_limit": radio_retransmission_limit,
        "frequency_hopping_mode": frequency_hopping_mode, "radio_flags": radio_flags,
        "registration_mode": rm, "broadcast_mode": broadcast_mode, "encryption_key": enc.hex(),
        "ptp_domain": ptp_domain, "debug_flags": debug_flags,
        "dscp": dscp, "ptp_mode": ptp_mode,
        # web GUI: ptpHybrid = ptpMode & 2, ptpSlaveOnly = ptpMode & 4
        "ptp_hybrid": bool(ptp_mode & 2), "ptp_slave_only": bool(ptp_mode & 4),
        "artist_net": ans,
        "bp_monitoring_threshold": bp_mon_threshold,
        "bp_monitoring_blocked_rssi": bp_mon_blocked_rssi,
        "bp_monitoring_interfered_rssi": bp_mon_interfered_rssi,
        "dect_scanner_blocked_rssi": dect_scanner_blocked_rssi,
        "dect_scanner_interfered_rssi": dect_scanner_interfered_rssi,
        "terms": terms,                               # JSON bpMonitoringList: beltpacks watched by BP monitoring
        "overrides": {
            "bluetooth": {"override": OVERRIDE_MODES.get(bt_mode, bt_mode), **bt},
            "brightness": {"override": OVERRIDE_MODES.get(br_mode, br_mode), "brightness_mode": br_sub,
                           "display": br1, "key": br2, "call_led_dim": br_i1, "status_led_dim": br_i2},
            "notification": {"override": OVERRIDE_MODES.get(notif_mode, notif_mode), "signalization": sig,
                             "signalization_raw": notif_sig, "beep_volume": notif_i1, "voice_volume": notif_i2},
            "speaker": {"override": OVERRIDE_MODES.get(spk_mode, spk_mode), "enabled": spk_enabled}
        }
    }

# ── section 1: partylines (radon::Partylines::unpackDataFromSaved / Partyline::unpack) ──
def up_partyline(s):
    """Partyline::unpack: i16 id, u8 (read and discarded by the firmware), name, +0x24 enabled,
    +0x28 timestamp, +0x2c showOnReply (JSON pl: id, name, enabled, type, showOnReply)."""
    pid=s.i16(); s.u8(); name=s.string(); en=s.boolean(); ts=s.u32(); sor=s.boolean()
    return {"id":pid,"name":name,"enabled":en,"timestamp":ts,"show_on_reply":sor}
def up_partylines_section(s):
    """[i16 removeCount][removeCount x (i16 id,u32 vc)][u8 count][count x Partyline]."""
    remove_count=s.i16()
    removals=[{"id":s.i16(),"timestamp":s.u32()} for _ in range(remove_count)]
    count=s.u8()
    return {"removals":removals,"partylines":[up_partyline(s) for _ in range(count)]}

# ── section 2: profiles (radon::Profiles::unpackDataFromSaved / Profiles::unpackDiffData) ──
def up_profiles_section(s):
    remove_count=s.i16()
    removals=[{"id":s.i16(),"timestamp":s.u32()} for _ in range(remove_count)]
    full_count=s.u8()
    profiles=[up_profile(s) for _ in range(full_count)]
    return {"removals":removals,"profiles":profiles}

# ── section 3: beltpacks (radon::RegisteredBPs::unpackDataFromSaved / RegisteredBPEntry::unpack) ──
def up_encryption_key(s): return s.read(0x20)
def up_registered_bp_entry(s):
    """RegisteredBPEntry::unpack: +4 TermId (the beltpack's id), +8 BPType, +0xc IPEI (u16 + u32),
    EncryptionKey(32), BPConfig, then +0x278 timestamp, +0x27c lastConnectTime, +0x280 config timestamp
    (setConfigTimestamp), +0x284 master timestamp (setMasterTimestamp). The IPEI is printed as 0x%04X %06X
    (IPEI::CharString) - the DECT identity on the beltpack's label."""
    a=s.u16(); b=s.u16(); c=s.u16(); d=s.u32()
    enc=up_encryption_key(s)
    cfg=up_bpconfig(s)
    tail=[s.u32() for _ in range(4)]
    return {"h0":a,"h1":b,"h2":c,"vc":d,"enc":enc,"config":cfg,"tail":tail,
            "term_id":a,"bp_type":b,"ipei":"0x%04X %06X" % (c, d),
            "timestamp":tail[0],"last_connect_time":tail[1],"config_timestamp":tail[2],"master_timestamp":tail[3]}
def up_beltpacks_section(s):
    count=s.u16()
    entries=[up_registered_bp_entry(s) for _ in range(count)]
    count2=s.u16()
    list2=[{"id":s.u16(),"timestamp":s.u32()} for _ in range(count2)]
    return {"beltpacks":entries,"list2":list2}

# ── section 4: priority nodes / antennas (radon::priorityNodes::PriorityNodes::unpackDataFromSaved) ──
def up_priority_nodes_section(s):
    """radon::priorityNodes::PriorityNodes::unpackAllData: [i16 count][count x (u32 node, PriorityNode)]."""
    count = s.i16()
    nodes = []
    for _ in range(count):
        node_id = s.u32()
        node_type = s.u8()
        user_id = s.i16()                              # JSON priorityNode userId
        name = s.string()
        label = s.string()
        n_terms = s.u16()
        terms = [s.u16() for _ in range(n_terms)]
        n_profs = s.u16()
        profs = [s.u16() for _ in range(n_profs)]
        ts = s.u32()
        nodes.append({
            "node_id": node_id, "node_type": node_type, "user_id": user_id,
            "name": name, "label": label, "terms": terms, "profiles": profs, "timestamp": ts
        })
    return {"nodes": nodes}

# ── section 5: audio filters (radon::AudioFilters::unpackDataFromSaved / AudioFilters::unpackAllData) ──
def up_audio_filter_section(s):
    # radon::AudioFilterSection::unpack = 12 bytes: u8 type, u16 freq, 8 bytes Q, i8 gain
    return {"type": s.u8(), "freq": s.u16(), "q": s.read(8).hex(), "gain": s.i8()}

def up_audio_filters_section(s):
    count1 = s.u16()
    vlist = [{"id": s.u8(), "timestamp": s.u32()} for _ in range(count1)]
    count2 = s.u8()
    filters = []
    for _ in range(count2):
        fid = s.u8(); u = s.u32(); name = s.string(); typ = s.u8()
        n_sec = s.u8()
        secs = [up_audio_filter_section(s) for _ in range(n_sec)]
        filters.append({"id": fid, "u": u, "name": name, "type": typ, "sections": secs})
    return {"version_list": vlist, "filters": filters}


# ── section 6: audio devices (radon::IODeviceConfig::unpackDataFromSaved / IODeviceConfig::unpackAllData) ──
def up_io_audio_channel_config(s):  # IODeviceAudioChannelConfig::unpack = 8 bytes
    return [s.i8(), s.u8(), s.i8(), s.i8(), s.u8(), s.i8(), s.u8(), s.i8()]

# One entry per connector pair of an I/O device (NSA-002A: XLR In N + XLR Out N). The 8 bytes above are
# [config id, mode, ?, ?, input type, input index, output type, index] - web GUI AudioChannelConfigPair
# {audioChannelConfigId, mode, input{type,index}, output{type,index}}. Indexes are 0-based: the GUI's NSA page
# titles pair i "Input i+1 / Output i+1". Mode names are the GUI's CONFIG.AUDIO_CHANNEL_MODES.
AUDIO_CHANNEL_MODES = {0: "Unused", 1: "4-Wire split", 2: "4-Wire", 3: "Input only", 4: "Output only"}

def connector_pairs(dev):
    out = []
    for i, c in enumerate(dev.get("channels", [])):
        out.append({"pair": i + 1, "config_id": c[0], "mode": c[1],
                    "mode_name": AUDIO_CHANNEL_MODES.get(c[1], f"?{c[1]}"),
                    "input": {"type": c[4], "index": c[5]}, "output": {"type": c[6], "index": c[7]}})
    return out

def plug_text(plug, devtype):
    """'XLR In 5' style name of one side of an audio channel (plug index is 0-based)."""
    if not plug or not plug.get("type") or plug.get("index", -1) < 0:
        return ""
    n = plug["index"] + 1
    if plug["type"] == 1 and devtype == 1:           # LegacyNSA plug on an NSA-002A: its XLR connectors
        return f"{n}"
    return f"{PLUG_NAMES.get(plug['type'], AUDIO_CHANNEL_TYPE.get(plug['type'], plug['type']))} {n}"

# plain names for the connector kinds (web GUI AUDIO_CHANNEL_TYPE)
PLUG_NAMES = {1: "NSA", 2: "Analog", 3: "Partyline", 9: "Universal (automatic)", 10: "Universal analog",
              11: "Digital left", 12: "Digital right", 13: "Digital left + right", 14: "USB audio left",
              15: "USB audio right", 16: "USB audio left + right", 17: "Headphone", 18: "2-Wire channel 1",
              19: "2-Wire channel 2"}

def channel_sides(ch):
    """[(side, connector number, disabled)] for the sides a channel uses ('in' / 'out').
    A channel is bound to its connector pair by channelIndex, and its flags say which sides it has. When a side
    of the pair is switched off on the device (the IODeviceAudioChannelConfig plug is type 0 / index -1) the
    channel's plug for that side is stored empty too, but the channel keeps its name and index: that side is a
    disabled port, not an unassigned one."""
    fl = ch.get("flags") or {}
    out = []
    for side, key, flag in (("in", "input_plug", "has_input"), ("out", "output_plug", "has_output")):
        plug = ch.get(key) or {}
        if plug.get("type") and plug.get("index", -1) >= 0:
            out.append((side, plug["index"] + 1, False))
        elif fl.get(flag) and ch.get("channel_index", -1) is not None and ch.get("channel_index", -1) >= 0:
            out.append((side, ch["channel_index"] + 1, True))
    return out

def channel_side_rows(ch, devtype):
    """[(connector, direction, disabled, sort key)] - one entry per side a channel uses, e.g. ('XLR 5', 'In', False).
    A 4-Wire channel gives an In and an Out entry, so both sides of an XLR pair can be listed next to each other."""
    rows = []
    for sd, n, dis in channel_sides(ch):
        if devtype == 1:
            conn = f"XLR {n}"
        else:
            plug = plug_text(ch.get("input_plug" if sd == "in" else "output_plug"), devtype)
            conn = plug or f"Channel {n}"
        rows.append((conn, "In" if sd == "in" else "Out", dis, (n, 0 if sd == "in" else 1)))
    return rows

def channel_connector_direction(ch, devtype):
    """(connector, direction, sort key) for one row per channel, e.g. ('XLR 4', 'In + Out (disabled)', ...)."""
    rows = channel_side_rows(ch, devtype)
    if not rows:
        return "No connector", "", (999, 9)
    conns = []
    for r in rows:
        if r[0] not in conns:
            conns.append(r[0])
    direction = " + ".join(d + (" (disabled)" if dis else "") for _, d, dis, _ in rows)
    return " / ".join(conns), direction, min(r[3] for r in rows)

def channel_connectors(ch, devtype):
    """Where an audio channel sits on its device, in XLR terms; a switched-off side is marked (disabled)."""
    sides = channel_sides(ch)
    if not sides:
        return "No connector"
    if devtype == 1:                                 # NSA-002A: XLR In N / XLR Out N
        return " + ".join(f"XLR {'In' if sd == 'in' else 'Out'} {n}" + (" (disabled)" if dis else "")
                          for sd, n, dis in sides)
    i = plug_text(ch.get("input_plug"), devtype); o = plug_text(ch.get("output_plug"), devtype)
    if i and i == o:                                 # e.g. a PunQtum partyline: one connection, both ways
        return i
    parts = []
    for sd, n, dis in sides:
        txt = (i if sd == "in" else o) or f"channel {n}"
        parts.append(f"{'In' if sd == 'in' else 'Out'}: {txt}" + (" (disabled)" if dis else ""))
    return " + ".join(parts)

# When each of an audio channel's function lists acts (web GUI AUDIO_CHANNEL_CONFIG_FUNCTIONS_DISPLAY_NAMES)
FUNCTION_WHEN = {"always_on": "Always", "on_talk": "On-Talk", "on_beep": "On-Notification/Beep",
                 "on_vox": "On-VOX"}

def channel_talk_listen(ch, pl, bp, chn, tr):
    """Chart of what an audio channel talks to and listens to, and when.
    Talk = the channel's input is sent to the destination; Listen = the destination is heard on its output.
    Talk+AlwaysListen talks under the list's condition and always listens. Other functions (e.g. Set Trigger)
    are returned as actions. Returns (rows, actions); a row is {destination, talk, listen, priority}."""
    rows = {}; actions = []
    lists = [(k, FUNCTION_WHEN[k], ch.get(k, [])) for k in ("always_on", "on_talk", "on_beep", "on_vox")]
    for t in ch.get("triggers", []):                 # trigger functions: act while a trigger input is set
        name = tr.get((t["port"]["type"], t["port"]["id"]), f"trigger {t['port']['id']}")
        lists.append(("trigger", f"On trigger '{name}'", [t["function"]]))
    for key, when, funcs in lists:
        for f in funcs:
            dest = resolve_target(f, pl, bp, chn, tr) or "(no destination)"
            fn = f["function"]
            talk = when if fn in (1, 2, 3) else ""
            listen = "Always" if fn == 2 else (when if fn in (3, 4) else "")
            if not talk and not listen:
                actions.append(f"{when}: {KEY_FUNCTIONS.get(fn, fn)} {dest}".strip())
                continue
            r = rows.setdefault(dest, {"destination": dest, "talk": [], "listen": [], "priority": set()})
            if talk and talk not in r["talk"]: r["talk"].append(talk)
            if listen and listen not in r["listen"]: r["listen"].append(listen)
            r["priority"].add(KEY_PRIORITIES.get(f.get("priority", 0), f.get("priority")))
    out = [{"destination": r["destination"], "talk": " / ".join(r["talk"]) or "-",
            "listen": " / ".join(r["listen"]) or "-", "priority": ", ".join(sorted(r["priority"]))}
           for r in rows.values()]
    return out, actions

TRIGGER_KINDS = {7: "GPI", 8: "GPO", 9: "Virtual"}   # AudioPortId types TriggerInput / TriggerOutput / TriggerVirtual

def gpio_chart(parsed):
    """One row per trigger pin: device, GPI/GPO, pin (1-based, as the web GUI's 'Input Pin {{$index + 1}}'),
    enabled, name, mode and every function assigned to it.
    Pins come from each device's inputTriggers / outputTriggers lists (1 = enabled); a trigger's port id is
    (ioDeviceConfig << 8) | pin. Assignments are collected from beltpack and profile keys, the always-on /
    on-talk / on-beep / on-VOX function lists of beltpacks, profiles and audio channels (e.g. Set Trigger to a
    GPO), and the trigger functions that fire while a GPI is active."""
    pl, bp, chn, tr = build_name_maps(parsed)
    trig = {(t["port"]["type"], t["port"]["id"]): t for t in parsed.get("gpio", {}).get("triggers", [])}
    uses = {}

    def add(port, text):
        uses.setdefault((port["type"], port["id"]), []).append(text)

    def scan(owner, cfg):
        for i, k in enumerate(cfg.get("keys", [])):
            if k["target"]["type"] in TRIGGER_KINDS:
                add(k["target"], f"{owner} Key {i + 1}: {k['function_name']} [{k['mode_name']}]")
        for lst, when in (("always_on", "Always-On"), ("on_talk", "On-Talk"), ("on_beep", "On-Notification/Beep"),
                          ("on_vox", "On-VOX")):
            for f in cfg.get(lst, []):
                if f["target"]["type"] in TRIGGER_KINDS:
                    add(f["target"], f"{owner} {when}: {f.get('function_name', f['function'])}")
        for t in cfg.get("triggers", []):
            f = t["function"]
            add(t["port"], f"while active -> {owner}: {f.get('function_name', f['function'])} "
                           f"{resolve_target(f, pl, bp, chn, tr)}".rstrip())

    for e in parsed.get("beltpacks", {}).get("beltpacks", []):
        scan(f"Beltpack {e['config'].get('name', '')}", e["config"])
    for p in parsed.get("profiles", {}).get("profiles", []):
        scan(f"Profile {p['name'] or p['config'].get('name', '')}", p["config"])
    for c in parsed.get("audio_channels", {}).get("channels", []):
        scan(f"Audio channel {c['channel'].get('name', '')}", c["channel"])

    rows, seen = [], set()
    for d in parsed.get("audio_devices", {}).get("devices", []):
        for typ, flags in ((7, d.get("trigger_in_flags", [])), (8, d.get("trigger_out_flags", []))):
            for pin, on in enumerate(flags):
                key = (typ, (d["config_id"] << 8) | pin)
                seen.add(key)
                t = trig.get(key, {})
                rows.append({"device": d["name"], "kind": TRIGGER_KINDS[typ], "pin": pin + 1, "enabled": bool(on),
                             "name": t.get("name", ""), "mode": t.get("mode_name", ""),
                             "timestamp": t.get("timestamp"), "assigned": uses.get(key, [])})
    for key, t in trig.items():                       # triggers not tied to a device pin (e.g. virtual)
        if key not in seen:
            rows.append({"device": "", "kind": TRIGGER_KINDS.get(key[0], key[0]), "pin": t.get("pin_index", 0) + 1,
                         "enabled": t.get("active", False), "name": t.get("name", ""), "mode": t.get("mode_name", ""),
                         "timestamp": t.get("timestamp"), "assigned": uses.get(key, [])})
    return rows


def beltpack_keys_on_channels(parsed):
    """{(port type, port id): 'Listen x12, Talk x2'} - beltpack keys (and reply) aimed at each audio channel.
    Input and output halves of a split share an id, so the port type is part of the key."""
    from collections import Counter
    use = {}
    for e in parsed.get("beltpacks", {}).get("beltpacks", []):
        for k in e["config"].get("keys", []):
            t = k["target"]
            if t["type"] in (3, 5, 6):
                use.setdefault((t["type"], t["id"]), Counter())[k["function_name"]] += 1
    return {p: ", ".join(f"{fn} x{n}" for fn, n in c.most_common()) for p, c in use.items()}

def channel_mode(ch, dev):
    """The mode of the connector pair a channel uses (e.g. '4-Wire split'), from its device's pair table."""
    pairs = connector_pairs(dev) if dev else []
    for side, key in (("input_plug", "input"), ("output_plug", "output")):
        plug = ch.get(side) or {}
        if not plug.get("type"):
            continue
        for pr in pairs:
            if pr[key]["type"] == plug["type"] and pr[key]["index"] == plug["index"]:
                return pr["mode_name"]
    idx = ch.get("channel_index", -1)                 # both sides switched off: the pair is still its index
    return pairs[idx]["mode_name"] if idx is not None and 0 <= idx < len(pairs) else ""

def up_nsa_config_data(s, devtype):
    flags = s.u8()
    n_in = 0; n_out = 0; n_pl = 0
    if devtype == 3: n_pl = 2
    elif devtype == 4: n_in = 4
    elif devtype == 5: n_out = 1
    elif devtype in (6, 10): n_in = 2; n_out = 2
    elif devtype == 7: n_in = 4; n_out = 2
    in_modes = [s.u8() for _ in range(n_in)]
    out_modes = [s.u8() for _ in range(n_out)]
    pl_modes = [s.u8() for _ in range(n_pl)]
    return {"flags": flags, "in_modes": in_modes, "out_modes": out_modes, "pl_modes": pl_modes}

def up_audio_devices_section(s):
    count = s.u8()
    devs = []
    for _ in range(count):
        cid = s.u8()
        devtype = s.u8()
        name1 = s.string(); name2 = s.string()      # JSON name, description
        r1 = s.read(4); u1 = s.u16()                # ioDeviceMulticastGroup (IP + port), twice
        r2 = s.read(4); u2 = s.u16()
        chan_cnt = s.u8()
        chans = [up_io_audio_channel_config(s) for _ in range(chan_cnt)]
        t1_cnt = s.u8(); t1 = list(s.read(t1_cnt))
        t2_cnt = s.u8(); t2 = list(s.read(t2_cnt))
        extra = None
        if devtype == 2:  # PunQtum
            pq_b1 = s.u8(); pq_c1 = s.u8(); pq_chans = []
            for _ in range(pq_c1):
                b = list(s.read(4)); s_a = s.string(); s_b = s.string()
                pq_chans.append({"bytes": b, "name1": s_a, "name2": s_b})
            pq_c2 = s.u8()
            pq_ctrl1 = [{"b": s.u8(), "name": s.string()} for _ in range(pq_c2)]
            pq_c3 = s.u8()
            pq_ctrl2 = [{"b": s.u8(), "name": s.string()} for _ in range(pq_c3)]
            extra = {"b1": pq_b1, "channels": pq_chans, "ctrl1": pq_ctrl1, "ctrl2": pq_ctrl2}
        elif devtype in (3, 4, 5, 6, 7, 10):
            extra = up_nsa_config_data(s, devtype)
        dev_id = s.read(8)
        vc = s.u32()
        devs.append({
            "config_id": cid, "devtype": devtype, "devtype_name": IO_DEVICE_TYPE.get(devtype, f"?{devtype}"),
            "name": name1, "name2": name2, "u1": u1, "u2": u2, "channels": chans,
            "description": name2,
            "multicast_groups": [f"{'.'.join(map(str, r1))}:{u1}", f"{'.'.join(map(str, r2))}:{u2}"],
            "trigger_in_flags": t1, "trigger_out_flags": t2,
            "extra": extra, "dev_id": dev_id.hex(), "timestamp": vc
        })
    return {"devices": devs}

# ── section 7: audio channels (radon::AudioChannels::unpackDataFromSaved / AudioChannels::unpackAllData) ──
def up_audio_channel_filter(s):
    """AudioChannelFilter (JSON lowPassFilter / highPassFilter): enabled, currentFrequency."""
    return {"enabled":s.boolean(),"frequency":s.u16()}
def up_audio_channel_limiter(s):
    """AudioChannelLimiter (JSON limiter): enabled, currentThreshold, currentAttackTime, currentReleaseTime."""
    return {"enabled":s.boolean(),"threshold":s.i8(),"attack_time":s.u16(),"release_time":s.u16()}

# AudioChannel flag word +0x44 (JsonSerializer::serializeAudioChannel / AudioChannel::dump)
CHANNEL_FLAG_BITS = {0x01: "enabled", 0x02: "has_input", 0x04: "has_output", 0x08: "input_mute",
                     0x10: "output_mute", 0x20: "phantom_power"}

def up_audio_channel(s):
    """AudioChannel::unpack. Names are serializeAudioChannel's JSON keys: ioDeviceConfig, channelIndex, the input
    and output stream index (+0xc / +0x10, 'Stream index' in AudioChannel::dump), input and output plug, name,
    flags, input currentGain (+0x48), output currentGain (+0x4c), output priorityDim (+0x50), audioPorts, the five
    function lists, vox (+0x54), showOnReply (+0xd0), then filters and limiters that exist only for some plugs."""
    first=s.u8()
    hdr=[s.i8(),s.i8(),s.i8(),s.u8(),s.i8(),s.u8(),s.i8()]
    f18=hdr[3]; f24=hdr[5]
    name=s.string()
    hdr2=[s.u8(),s.i8(),s.i8(),s.i8()]
    ports=up_audio_ports_list(s)
    fA=[up_single_function(s) for _ in range(s.u8())]
    trig=[up_trigger_function(s) for _ in range(s.u8())]
    fC=[up_single_function(s) for _ in range(s.u8())]
    fD=[up_single_function(s) for _ in range(s.u8())]
    fE=[up_single_function(s) for _ in range(s.u8())]
    vox=up_voxbase(s); show_on_reply=s.boolean()
    filters=[]; limiters=[]
    if f18==10 or f18 in (18,19): filters.append(up_audio_channel_filter(s))
    if f18 in (18,19):            filters.append(up_audio_channel_filter(s))
    if f24 in (18,19):            filters.append(up_audio_channel_filter(s)); filters.append(up_audio_channel_filter(s))
    if f18==10 or f18 in (18,19): limiters.append(up_audio_channel_limiter(s))
    if f24 in (17,18,19):         limiters.append(up_audio_channel_limiter(s))
    # AudioChannel::unpack: +0x18/+0x1c input plug (type, index), +0x24/+0x28 output plug (type, index)
    # (setInputPlug / setOutputPlug; web GUI channel.input / channel.output). Type 0 / index -1 = side not used,
    # so an output-only channel has input type 0. 'type' is the plug type of whichever side is in use.
    typ = f18 or f24
    return {"type":typ,"type_name":AUDIO_CHANNEL_TYPE.get(typ,f"?{typ}") if typ else "Disabled","name":name,"ports":ports,
            "input_plug":{"type":f18,"type_name":AUDIO_CHANNEL_TYPE.get(f18,f"?{f18}") if f18 else "None","index":hdr[4]},
            "output_plug":{"type":f24,"type_name":AUDIO_CHANNEL_TYPE.get(f24,f"?{f24}") if f24 else "None","index":hdr[6]},
            "device_config_id":first,"channel_index":hdr[0],
            "input_stream_index":hdr[1],"output_stream_index":hdr[2],
            "flags":{n: bool(hdr2[0] & b) for b, n in CHANNEL_FLAG_BITS.items()},
            "input_gain":hdr2[1],"output_gain":hdr2[2],"output_priority_dim":hdr2[3],
            "show_on_reply":show_on_reply,
            "triggers":trig,"vox":vox,"filters":filters,"limiters":limiters,
            # the five lists in file order are the web GUI's audio channel functions: alwaysOnFunctions,
            # triggerFunctions, onCallFunctions ('On-Talk'), onSilentCallFunctions ('On-Notification/Beep'),
            # onVoxFunctions ('On-VOX') (AUDIO_CHANNEL_CONFIG_FUNCTIONS); they match the usage bits in each
            # partyline entry's mask (0x02 always-on, 0x08 trigger, 0x10 on-talk, 0x20 on-beep, 0x40 on-VOX)
            "always_on":fA,"on_talk":fC,"on_beep":fD,"on_vox":fE}
def up_audio_channels_section(s):
    count1=s.i16(); vlist=[{"id":s.u8(),"timestamp":s.u32()} for _ in range(count1)]
    count2=s.i16()
    chans=[{"port":up_audio_port_id(s),"timestamp":s.u32(),"channel":up_audio_channel(s)}
           for _ in range(count2)]
    return {"version_list":vlist,"channels":chans}

# ── section 8: gpio / triggers (radon::TriggerConfig::unpackDataFromSaved / TriggerConfig::unpackAllData) ──
def up_gpio_section(s):
    count1 = s.i16()
    diffs = [{"id": s.u8(), "timestamp": s.u32()} for _ in range(count1)]
    count2 = s.i16()
    triggers = []
    for _ in range(count2):
        port = up_audio_port_id(s)
        vc = s.u32()
        dev_cfg_id = s.u8()
        pin_idx = s.u8()
        active = s.boolean()
        trig_name = s.string()
        trig_mode = s.u8()
        triggers.append({
            "port": port, "timestamp": vc,
            "device_config_id": dev_cfg_id, "pin_index": pin_idx, "active": active,
            "name": trig_name, "mode": trig_mode,
            "mode_name": TRIGGER_MODES.get(trig_mode, f"?{trig_mode}")
        })
    return {"diffs": diffs, "triggers": triggers}



# ── container framing (firmware-confirmed) ─────────────
BOLERO_SECTIONS = [
    "network", "partylines", "profiles", "beltpacks", "antennas",
    "unknown", "audio_devices", "audio_channels", "gpio"
]

def read_container(raw):
    data = zlib.decompress(raw[7:])
    count = data[2]; p = 23; secs = []
    for i in range(count):
        version, length = struct.unpack_from(">HI", data, p + 8)
        start = p + 14; end = start + length
        secs.append({"index": i, "role": BOLERO_SECTIONS[i] if i < len(BOLERO_SECTIONS) else f"s{i}",
                     "version": version, "start": start, "end": end, "length": length})
        p = end
    return data, secs


SECTION_READERS = {
    "network": up_net_settings_section,
    "partylines": up_partylines_section,
    "profiles": up_profiles_section,
    "beltpacks": up_beltpacks_section,
    "antennas": up_priority_nodes_section,
    "unknown": up_audio_filters_section,
    "audio_devices": up_audio_devices_section,
    "audio_channels": up_audio_channels_section,
    "gpio": up_gpio_section,
}

def parse_file(path):
    """Parse a .bol; returns {role: parsed} for all 9 sections, each gated on exact byte consumption."""
    data, secs = read_container(open(path, "rb").read())
    out = {"_sections": secs}
    for sec in secs:
        r = SECTION_READERS.get(sec["role"])
        if not r: continue
        s = InputStream(data, sec["start"], sec["end"]); s.p = sec["start"]
        parsed = r(s)
        if s.p != sec["end"]:
            raise ValueError(f"{sec['role']}: consumed to {s.p}, expected {sec['end']} "
                             f"(residual {sec['end']-s.p})")
        out[sec["role"]] = parsed
    return out


# ── target name resolution ──────────────────────────────────────────────────
def pin_text(v, admin=None):
    """A network PIN as the Bolero web GUI means it (3.4.1 app-*.js: PINS, PIN_LENGTH.ADMIN = 4).
    Stored as the int16 of parseInt(4-digit code), so leading zeros are gone: 0 is '0000', 12 is '0012'.
    -1 PIN_CODE_NONE, -2 PIN_CODE_USE_ADMIN (OTA only: 'Use Admin PIN for OTA Registration'),
    -100 PIN_CODE_INVALID (an old admin PIN the firmware replaces, JsonSerializer::updateOldAdminPin)."""
    if v is None or v == "":
        return ""
    if v == -1:
        return "No PIN"
    if v == -2:
        return "Uses Admin PIN" + (" (%s)" % pin_text(admin) if admin is not None and admin >= 0 else "")
    if v == -100:
        return "Invalid (old PIN, reset by firmware)"
    return "%04d" % v if v >= 0 else "Unknown (%d)" % v


# radio settings as the 3.4.1 web GUI names them (NET_SETTINGS.NWS_RADIO_POWERS, RADIORETRANSMIT_LIMIT) and
# radioFlags bits (NetSettingsData::unpack). The DECT country / region is a per-device setting changed in the
# Service view (6-digit Service PIN); it is not part of the saved network settings.
RADIO_POWER_NAMES = {0: "Normal", 1: "Low", 2: "Ultra Low"}
RETRANSMIT_NAMES = {1: "Low", 2: "Medium", 3: "High", 4: "Very High"}
RADIO_FLAG_NAMES = ["Radio enabled", "Radio priority", "BP monitoring threshold", "High 2.4 GHz radio power",
                    "DECT scanner", "Web server encryption"]
PIN_NOTE = ("The Admin PIN logs in to the beltpack admin menu and to the web GUI as Admin (4 digits). "
            "The 6-digit Service PIN, needed to change the DECT region, is set per device, is not in show files, "
            "and is issued by Riedel support.")

def radio_flags_text(flags):
    return ", ".join(n for i, n in enumerate(RADIO_FLAG_NAMES) if (flags or 0) >> i & 1) or "None"

def named(table, v):
    return "%s (%s)" % (table[v], v) if v in table else v


def ts_text(v):
    """The u32 change markers in a .bol are Unix times (UTC)."""
    import datetime
    try:
        return datetime.datetime.fromtimestamp(v, datetime.timezone.utc).strftime("%Y-%m-%d %H:%M") if v else ""
    except (OverflowError, OSError, ValueError, TypeError):
        return str(v)


def build_name_maps(parsed):
    pl = {p["id"]: p["name"] for p in parsed.get("partylines", {}).get("partylines", [])}
    bp = {e["h0"]: e["config"]["name"] for e in parsed.get("beltpacks", {}).get("beltpacks", [])}
    # audio channels and triggers are keyed by (port type, id): the input and output halves of a 4-Wire split
    # share an id (e.g. 516 = the input half and the output half of one XLR pair), as can trigger inputs and outputs
    ch = {(c["port"]["type"], c["port"]["id"]): c["channel"]["name"]
          for c in parsed.get("audio_channels", {}).get("channels", []) if c["port"]["id"] is not None}
    tr = {(t["port"]["type"], t["port"]["id"]): t["name"] for t in parsed.get("gpio", {}).get("triggers", [])
          if t["port"]["id"] is not None}
    return pl, bp, ch, tr

def resolve_target(k, pl, bp, ch, tr):
    t = k["target"]; typ = t["type"]; tid = t["id"]
    if typ == 0: return ""
    if typ == 2: return f"PL: {pl.get(tid, 'Partyline ' + str(tid))}"
    if typ == 1: return f"P2P: {bp.get(tid, 'BP ' + str(tid))}"
    if typ in (3, 5, 6):
        nm = ch.get((typ, tid)); d = "BiDir" if typ == 3 else ("Listen" if typ == 5 else "Talk")
        return f"Audio {d}: {nm}" if nm else f"Audio {d} (ch {tid})"
    if typ in (7, 8, 9):
        nm = tr.get((typ, tid))
        return f"Trigger: {nm}" if nm else f"Trigger {tid}"
    return f"{t['type_name']} {tid}"


# ── workbook styling (self-contained; formerly borrowed from art_to_excel) ─────────────────────────────
class _XLStyle:
    """Title banner, header row, cell writer and column widths for the Bolero workbook."""
    C = {'navy': '1F3864', 'blue': '2E5FA3', 'slate': '4A4A4A', 'border': 'D9D9D9',
         'row_alt': 'F4F7FC', 'row_white': 'FFFFFF'}

    @staticmethod
    def _clean(v):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        return "" if v is None else ILLEGAL_CHARACTERS_RE.sub("", str(v))

    @staticmethod
    def _fill(hex_c):
        from openpyxl.styles import PatternFill
        return PatternFill(start_color=hex_c, end_color=hex_c, fill_type="solid")

    @staticmethod
    def _border():
        from openpyxl.styles import Border, Side
        s = Side(style="thin", color=_XLStyle.C['border'])
        return Border(left=s, right=s, top=s, bottom=s)

    @staticmethod
    def title_banner(ws, title, subtitle, max_col=8, show_home_link=True):
        from openpyxl.styles import Alignment, Font
        S = _XLStyle
        last = max_col - 1 if (show_home_link and max_col > 2) else max_col
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
        c1 = ws.cell(row=1, column=1, value=S._clean(title))
        c1.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        c1.fill = S._fill(S.C['navy'])
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if last != max_col:
            home = ws.cell(row=1, column=max_col, value="<- Summary")
            home.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF", underline="single")
            home.fill = S._fill(S.C['navy'])
            home.alignment = Alignment(horizontal="center", vertical="center")
            home.hyperlink = "#'Summary'!A1"
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        c2 = ws.cell(row=2, column=1, value=S._clean(subtitle))
        c2.font = Font(name="Calibri", size=9, italic=True, color="FFFFFF")
        c2.fill = S._fill(S.C['blue'])
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 18

    @staticmethod
    def header_row(ws, row_idx, headers, bg=None, fg="FFFFFF"):
        from openpyxl.styles import Alignment, Font
        S = _XLStyle
        ws.row_dimensions[row_idx].height = 24
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=S._clean(h))
            c.font = Font(name="Calibri", size=10, bold=True, color=fg)
            c.fill = S._fill(bg or S.C['slate'])
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = S._border()

    @staticmethod
    def write_cell(ws, row_idx, col_idx, val, bg=None, fg="000000", bold=False, align="left", size=9):
        from openpyxl.styles import Alignment, Font
        S = _XLStyle
        c = ws.cell(row=row_idx, column=col_idx, value=S._clean(val))
        c.font = Font(name="Calibri", size=size, bold=bold, color=fg)
        if bg:
            c.fill = S._fill(bg)
        c.alignment = Alignment(horizontal=align, vertical="center")
        c.border = S._border()
        return c

    @staticmethod
    def auto_width(ws, extra=3, max_w=65):
        from openpyxl.utils import get_column_letter
        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w + extra, 10), max_w)


def json_default(o):
    """JSON for values json can't write: raw bytes (e.g. encryption keys) as hex, anything else as text."""
    return o.hex() if isinstance(o, (bytes, bytearray)) else str(o)


def to_json_dict(parsed, filename=""):
    """The Bolero JSON export: every section with its decoded, named fields."""
    return {"format": "Riedel Bolero Standalone .bol", "file": filename, **parsed}


def to_json_text(parsed, filename=""):
    import json
    return json.dumps(to_json_dict(parsed, filename), indent=2, default=json_default)


def export_to_excel(path, out_xlsx=None):
    """Faithful .bol -> formatted .xlsx using all 9 fully validated sections."""
    import openpyxl
    A2E = _XLStyle                  # styling helpers (no dependency on the Artist exporter)
    import pathlib
    parsed = parse_file(path)
    pl, bp, ch, tr = build_name_maps(parsed)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    filename = pathlib.Path(path).name
    net = parsed.get("network", {})
    sys_name = net.get("show_name", "Bolero Standalone")

    # --- Summary Sheet ---
    ws_sum = wb.create_sheet("Summary")
    A2E.title_banner(ws_sum, f"Bolero Standalone Summary: {sys_name}", f"File: {filename}", max_col=3, show_home_link=False)
    A2E.header_row(ws_sum, 4, ["Contents", "Description", "Items"], bg=A2E.C['navy'])
    
    pl_len = len(parsed.get("partylines", {}).get("partylines", []))
    ch_len = len(parsed.get("audio_channels", {}).get("channels", []))
    prof_len = len(parsed.get("profiles", {}).get("profiles", []))
    bp_len = len(parsed.get("beltpacks", {}).get("beltpacks", []))
    ant_len = len(parsed.get("antennas", {}).get("nodes", []))
    dev_len = len(parsed.get("audio_devices", {}).get("devices", []))
    gpio_len = len(parsed.get("gpio", {}).get("triggers", []))

    toc = [
        ("Network", "Show & Network Settings", 1),
        ("Partylines", "Configured Partylines", pl_len),
        ("Audio Channels", "Configured Audio Channels", ch_len),
        ("Profiles", "Beltpack Profiles", prof_len),
        ("Beltpacks", "Registered Beltpacks", bp_len),
        ("Antenna Priority", "Antennas with priority beltpacks/profiles (antennas are not stored otherwise)", ant_len),
        ("Audio Devices", "NSA & PunQtum Devices", dev_len),
        ("GPIO Triggers", "GPIO Triggers", gpio_len),
    ]

    for r, (sheet_name, desc, count) in enumerate(toc, 5):
        c1 = A2E.write_cell(ws_sum, r, 1, sheet_name, bold=True)
        c1.hyperlink = f"#'{sheet_name}'!A1"
        c1.font = openpyxl.styles.Font(name="Calibri", size=10, bold=True, color="0563C1", underline="single")
        A2E.write_cell(ws_sum, r, 2, desc)
        A2E.write_cell(ws_sum, r, 3, count, align="center")

    def sheet(title, headers, subtitle=""):
        ws = wb.create_sheet(title)
        A2E.title_banner(ws, title, f"Bolero Standalone | {subtitle}" if subtitle else "Bolero Standalone", max_col=max(3, len(headers)), show_home_link=True)
        A2E.header_row(ws, 4, headers)
        return ws

    def append_row(ws, r_idx, row_data):
        for c_idx, val in enumerate(row_data, 1):
            A2E.write_cell(ws, r_idx, c_idx, val)

    # 1. Show & Network
    ws = sheet("Network", ["Setting", "Value"], "Show & Network Configuration")
    ipn = net.get("artist_net", {})
    rm = net.get("registration_mode", {})
    rows = [
        ("Show Name", net.get("show_name", "")),
        ("Network Label", net.get("label", "")),
        ("System Mode", net.get("system_mode", "")),
        ("— PINs —", ""),
        ("Admin PIN (beltpack admin menu + web GUI Admin login)", pin_text(net.get("admin_pin"))),
        ("OTA Registration PIN", pin_text(net.get("ota_pin"), net.get("admin_pin"))),
        ("Service PIN (6 digits, DECT region changes)", "Not in show files - issued by Riedel support"),
        ("Note", PIN_NOTE),
        ("— DECT / Radio —", ""),
        ("DECT region", "Set per device in the Service view - not saved in the show file"),
        ("Radio Power (DECT)", named(RADIO_POWER_NAMES, net.get("radio_power", ""))),
        ("Radio Power (2.4 GHz)", named(RADIO_POWER_NAMES, net.get("radio_power_2g4", ""))),
        ("Retransmit Level", named(RETRANSMIT_NAMES, net.get("radio_retransmission_limit", ""))),
        ("Frequency Hopping Mode (0-15)", net.get("frequency_hopping_mode", "")),
        ("Radio options on", radio_flags_text(net.get("radio_flags"))),
        ("Radio Flags (raw)", net.get("radio_flags", "")),
        ("Broadcast Mode", net.get("broadcast_mode", "")),
        ("— BP / DECT Monitoring —", ""),
        ("BP Monitoring Threshold", net.get("bp_monitoring_threshold", "")),
        ("BP Monitoring Blocked RSSI", net.get("bp_monitoring_blocked_rssi", "")),
        ("BP Monitoring Interfered RSSI", net.get("bp_monitoring_interfered_rssi", "")),
        ("DECT Scanner Blocked RSSI", net.get("dect_scanner_blocked_rssi", "")),
        ("DECT Scanner Interfered RSSI", net.get("dect_scanner_interfered_rssi", "")),
        ("— Timing / PTP —", ""),
        ("PTP Mode (raw)", net.get("ptp_mode", "")),
        ("PTP Hybrid", "Yes" if net.get("ptp_hybrid") else "No"),
        ("PTP Slave Only", "Yes" if net.get("ptp_slave_only") else "No"),
        ("PTP Domain", net.get("ptp_domain", "")),
        ("Time Source", {0: "Internal", 1: "NTP", 2: "PTP"}.get(net.get("time_source"), net.get("time_source", ""))),
        ("— Registration —", ""),
        ("Registration", "Enabled" if rm.get("registration_enabled") else "Disabled"),
        ("  OTA registration", "On" if rm.get("ota_enabled") else "Off"),
        ("  NFC registration", "On" if rm.get("nfc_enabled") else "Off"),
        ("  Charger registration", "On" if rm.get("charger_enabled") else "Off"),
        ("  Registration timeout", rm.get("timeout", "")),
        ("  Registration profile ID", rm.get("profile_id", "")),
        ("— Network / Security —", ""),
        ("Broadcast Encryption Key", net.get("encryption_key", "")),
        # DSCPSettings = JSON dscpPtp, dscpRtp, dscpControl (web GUI defaults 46 / 34 / 36)
        ("DSCP (PTP / Audio RTP / Control)", " / ".join(str(x) for x in net.get("dscp", [])) or "—"),
        ("Artist Network ID", ipn.get("network_id", "")),
        ("Artist Cluster ID", ipn.get("cluster_id", "")),
        ("Artist Discovery Multicast IP", ipn.get("mc_announce_ip", "")),
        ("Artist Discovery Port", ipn.get("mc_announce_port", "")),
        ("Multicast TTL", net.get("multicast_ttl", "")),
    ]
    _ants = parsed.get("antennas", {}).get("nodes", [])
    _ants_with_prio = [a for a in _ants if a.get("terms") or a.get("profiles")]
    if _ants:
        rows.append(("— Licensed Features (inferred from config) —", ""))
        rows.append(("Beltpack Priority per Antenna",
                     f"CONFIGURED on {len(_ants_with_prio)}/{len(_ants)} antennas "
                     f"(implies P2/priority license in net space)" if _ants_with_prio
                     else "antennas present, no priority lists set"))
    for r_idx, (k, v) in enumerate(rows, 5):
        A2E.write_cell(ws, r_idx, 1, k, bold=not bool(v) and k.startswith('—'))
        A2E.write_cell(ws, r_idx, 2, v)

    # 2. Partylines
    ws = sheet("Partylines", ["ID", "Name"])
    for r_idx, p in enumerate(parsed.get("partylines", {}).get("partylines", []), 5):
        append_row(ws, r_idx, [p["id"], p["name"]])

    # 3. Audio Channels
    # one row per connector side (a 4-Wire channel gives an In and an Out row), sorted by device, connector,
    # In before Out - so XLR 5 In and XLR 5 Out sit next to each other
    ws = sheet("Audio Channels", ["Device", "Connector", "Direction", "Status", "Channel", "Connector mode",
                                  "Partylines", "VOX", "Port ID"])
    devs_by_id = {d["config_id"]: d for d in parsed.get("audio_devices", {}).get("devices", [])}
    side_rows = []
    for c in parsed.get("audio_channels", {}).get("channels", []):
        chn = c["channel"]; d = devs_by_id.get(chn.get("device_config_id"), {})
        pls = ", ".join(pl.get(p["port"]["id"], f"PL {p['port']['id']}") for p in chn.get("ports", [])
                        if p["port"]["type"] == 2)
        for conn, direction, dis, key in channel_side_rows(chn, d.get("devtype")) or [("No connector", "", False, (999, 9))]:
            side_rows.append(((d.get("name", ""), key), [d.get("name", ""), conn, direction,
                              "Disabled" if dis else "Enabled", chn["name"], channel_mode(chn, d), pls,
                              vox_text(chn.get("vox")), c["port"]["id"]]))
    for r_idx, (_, row) in enumerate(sorted(side_rows, key=lambda x: x[0]), 5):
        append_row(ws, r_idx, row)

    # 3b. Talk / listen chart: one row per channel and destination (idle channels get one row too)
    ws = sheet("Talk-Listen", ["Device", "Connector", "Direction", "Channel", "Destination", "Talk", "Listen",
                               "Priority", "Other actions", "Beltpack keys on it", "VOX"],
               "Talk = the channel's input is sent to the destination; Listen = the destination is heard on its output")
    r_idx = 5
    chans_sorted = sorted(parsed.get("audio_channels", {}).get("channels", []),
                          key=lambda c: (devs_by_id.get(c["channel"].get("device_config_id"), {}).get("name", ""),
                                         channel_connector_direction(c["channel"], devs_by_id.get(c["channel"].get("device_config_id"), {}).get("devtype"))[2]))
    bp_keys = beltpack_keys_on_channels(parsed)
    for c in chans_sorted:
        chn = c["channel"]; d = devs_by_id.get(chn.get("device_config_id"), {})
        rows, actions = channel_talk_listen(chn, pl, bp, ch, tr)
        rows = rows or [{"destination": "(nothing assigned)", "talk": "-", "listen": "-", "priority": ""}]
        for i, r in enumerate(rows):
            conn, direction, _ = channel_connector_direction(chn, d.get("devtype"))
            append_row(ws, r_idx, [d.get("name", ""), conn, direction, chn["name"],
                                   r["destination"], r["talk"], r["listen"], r["priority"],
                                   "; ".join(actions) if i == 0 else "",
                                   bp_keys.get((c["port"]["type"], c["port"]["id"]), "") if i == 0 else "",
                                   vox_text(chn.get("vox")) if i == 0 else ""])
            r_idx += 1

    # Keys as columns: one row per profile / beltpack, and Function / Mode / Target for each of the 7 keys
    # (keys 1-7). In Standalone the 7th key is a normal key (Talk, Listen, Reply, ...); only in Artist-integrated
    # mode is it fixed to Reply. Unused keys stay blank; a non-standard priority is shown after the function.
    prof_by_id_all = {pr["config"]["id"]: (pr["name"] or pr["config"]["name"])
                      for pr in parsed.get("profiles", {}).get("profiles", [])}
    key_labels = [f"Key {i}" for i in range(1, 8)]
    key_headers = [f"{lbl} {part}" for lbl in key_labels for part in ("Function", "Mode", "Target")]

    def key_cells(cfg):
        cells = []
        for k in cfg.get("keys", []):
            if k["function"] == 0:
                cells += ["", "", ""]
                continue
            fn = k["function_name"] + (f" [{k['priority_name']} priority]" if k.get("priority") else "")
            cells += [fn, k["mode_name"], resolve_target(k, pl, bp, ch, tr) or "(no destination)"]
        return cells

    # 4. Profiles
    users_of = {}
    for e in parsed.get("beltpacks", {}).get("beltpacks", []):
        users_of[e["config"].get("id")] = users_of.get(e["config"].get("id"), 0) + 1
    ws = sheet("Profiles", ["Profile ID", "Profile", "Beltpacks using it"] + key_headers)
    for r_idx, pr in enumerate(parsed.get("profiles", {}).get("profiles", []), 5):
        pid = pr["config"]["id"]
        append_row(ws, r_idx, [pid, pr["name"], users_of.get(pid, 0)] + key_cells(pr["config"]))

    # 5. Beltpacks
    ws = sheet("Beltpacks", ["User ID", "Beltpack", "Profile"] + key_headers)
    for r_idx, e in enumerate(sorted(parsed.get("beltpacks", {}).get("beltpacks", []),
                                     key=lambda e: (e["config"].get("bp_number") or 0, e["h0"])), 5):
        cfg = e["config"]
        append_row(ws, r_idx, [cfg.get("bp_number"), cfg["name"], prof_by_id_all.get(cfg.get("id"), "")]
                   + key_cells(cfg))

    # 5b. Beltpack info: identity and the decoded per-beltpack settings (one row per beltpack)
    band = {3300: "1.9 GHz", 3303: "1.9 GHz", 3304: "2.4 GHz", 3306: "2.4 GHz"}
    ws = sheet("Beltpack Info", ["User ID", "Name", "IPEI", "Type code", "Band", "TermId (internal)", "Profile ID",
                                 "Last connected (UTC)", "Config changed (UTC)", "Headset vol", "Speaker vol",
                                 "Sidetone", "Headset mic gain", "Internal mic gain", "Mic type", "Language",
                                 "VOX", "Silent mode", "Speaker enabled", "Echo cancellation", "Noise filter",
                                 "Call signal", "Silent call signal", "Low battery signal", "Out of range signal",
                                 "Default priority", "Priority exceptions"])
    for r_idx, e in enumerate(parsed.get("beltpacks", {}).get("beltpacks", []), 5):
        c = e["config"]; v = c.get("volumes", {}); fl = c.get("flags", {}); sg = c.get("signalization", {})
        pr = c.get("priority", {})
        append_row(ws, r_idx, [c.get("bp_number", ""), c.get("name", ""), e.get("ipei", ""), e["h1"], band.get(e["h1"], ""),
                               e["h0"], c.get("id", ""), ts_text(e.get("last_connect_time")),
                               ts_text(e.get("config_timestamp")), v.get("headset_volume"), v.get("speaker_volume"),
                               v.get("sidetone_volume"), v.get("headset_mic_input_gain"),
                               v.get("internal_mic_input_gain"), c.get("mic_type_name"), c.get("language_name"),
                               vox_text(c.get("vox")), fl.get("silent_mode"), fl.get("speaker_enable"),
                               fl.get("echo_cancellation"), c.get("noise_filter"), sg.get("call"),
                               sg.get("silent_call"), sg.get("low_battery"), sg.get("out_of_range"),
                               pr.get("default_priority_name", ""), len(pr.get("exceptions", []))])

    # 6. Antennas
    import datetime as _dt
    # firmware radon::isDectAntenna(type 1) / is2G4Antenna(type 4); other node types are not antennas
    NODE_TYPE_NAMES = {1: "DECT antenna (1.9 GHz)", 4: "2.4 GHz antenna"}
    bp_by_h0 = {e["h0"]: e["config"]["name"] for e in parsed.get("beltpacks", {}).get("beltpacks", [])}
    prof_by_id = {pr["config"]["id"]: (pr["name"] or pr["config"]["name"])
                  for pr in parsed.get("profiles", {}).get("profiles", [])}
    ws = sheet("Antenna Priority", ["Antenna Name", "Type", "User ID", "Antenna ID",
                            "Priority Beltpacks", "Priority Profiles", "Last Modified"])
    for r_idx, n in enumerate(parsed.get("antennas", {}).get("nodes", []), 5):
        members = ", ".join(bp_by_h0.get(t, f"#{t}") for t in n["terms"])
        profs = ", ".join(prof_by_id.get(p, f"#{p}") for p in n["profiles"])
        try: ts = _dt.datetime.utcfromtimestamp(n["timestamp"]).strftime("%Y-%m-%d %H:%M")
        except Exception: ts = n["timestamp"]
        append_row(ws, r_idx, [n["name"], NODE_TYPE_NAMES.get(n["node_type"], f"Node type {n['node_type']}"),
                               n["user_id"], hex(n["node_id"]), members, profs, ts])

    # 7. Audio Devices
    ws = sheet("Audio Devices", ["Config ID", "Device Type", "Name", "Device ID", "Connector pairs", "Connector modes"])
    for r_idx, d in enumerate(parsed.get("audio_devices", {}).get("devices", []), 5):
        append_row(ws, r_idx, [d["config_id"], d["devtype_name"], d["name"], d["dev_id"], len(d["channels"]),
                               "; ".join(f"{'XLR' if d['devtype'] == 1 else 'Channel'} {pr['pair']}: {pr['mode_name']}"
                                         for pr in connector_pairs(d)) or "No connectors set up"])

    # 8. GPIO Triggers
    # one row per GPI / GPO pin of each device, with every function assigned to it
    ws = sheet("GPIO Triggers", ["Device", "GPI / GPO", "Pin", "Enabled", "Name", "Mode", "Assigned functions",
                                 "Last Changed"],
               "GPI = trigger input (fires trigger functions); GPO = trigger output (driven by Set Trigger functions)")
    for r_idx, g in enumerate(gpio_chart(parsed), 5):
        append_row(ws, r_idx, [g["device"], g["kind"], g["pin"], "Yes" if g["enabled"] else "No", g["name"],
                               g["mode"], "; ".join(g["assigned"]) or "-", ts_text(g["timestamp"])])

    # Auto-adjust column widths & add frozen panes
    for ws in wb.worksheets:
        A2E.auto_width(ws)
        # keep the name columns in view while scrolling across the key columns
        ws.freeze_panes = "D5" if ws.title in ("Beltpacks", "Profiles") else "A5"

    out_xlsx = out_xlsx or path.rsplit(".", 1)[0] + "_faithful.xlsx"
    wb.save(out_xlsx)
    return out_xlsx

if __name__ == "__main__":
    files = sys.argv[1:]
    if not files:
        sys.exit("usage: python bol_faithful.py FILE.bol [FILE.bol ...]   (checks that every section parses exactly)")
    all_ok = True
    for f in files:
        name = f.replace("\\", "/").split("/")[-1]
        try:
            r = parse_file(f)
            counts = (f"net='{r['network']['show_name'][:14]}' "
                      f"pl={len(r['partylines']['partylines'])} "
                      f"prof={len(r['profiles']['profiles'])} "
                      f"bp={len(r['beltpacks']['beltpacks'])} "
                      f"ant={len(r['antennas']['nodes'])} "
                      f"devs={len(r['audio_devices']['devices'])} "
                      f"ch={len(r['audio_channels']['channels'])} "
                      f"trig={len(r['gpio']['triggers'])}")
            print(f"[OK] {name[:34]:34s} {counts}")
        except Exception as e:
            all_ok = False; print(f"[FAIL] {name[:34]:34s} {type(e).__name__}: {e}")
    print("\nALL 9 SECTIONS FULLY VALIDATED AND EXACT-CONSUMED ON ALL FILES:", all_ok)
